"""
Spike Fase 0 — akurasi, waktu, dan biaya VLM pada resi asli (PRD §13 "Spike Fase 0").

    # 1. Pecah PDF bundel scan menjadi gambar per halaman (300 dpi):
    uv run python -m csm_ml.spike_vlm pdf ~/csm-data/"Mr. Said Cont 257.pdf" --out ~/csm-data/spike/img

    # 2. Jalankan model pada 50 gambar, bandingkan dengan workbook final (sheet Data Laut):
    uv run python -m csm_ml.spike_vlm run --images ~/csm-data/spike/img \\
        --labels ~/csm-data/"Container 257.xlsx" --container 257 --n 50 --out ~/csm-data/spike/hasil \\
        --price-in 0.60 --price-out 3.60        # USD per 1 juta token, dari halaman harga penyedia

    # Uji alat tanpa HF_TOKEN (klien tiruan yang menjawab dari label):
    uv run python -m csm_ml.spike_vlm run ... --fake

Data resi berisi paspor & HP: semua input dan output harus di luar repo. Laporan di layar hanya angka.
"""

from __future__ import annotations

import argparse
import json
import re
import statistics
import time
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from rapidfuzz.distance import Levenshtein

from .config import load_settings
from .hf_client import image_data_url, make_client
from .normalize import normalize_numeric, normalize_passport, normalize_phones, parse_count, parse_weight_kg

INSTRUCTION = "Baca resi ini dan isi semua field sesuai schema."

# Kolom sheet `Data Laut Container-{no}` (PRD §4).
DATA_LAUT_COLUMNS = {
    "sender_name": "B",
    "recipient_name": "C",
    "address": "D",
    "serial_no": "F",
    "passport_no": "G",
    "recipient_phone": "H",
    "tas": "I",
    "krt": "J",
    "krg": "K",
    "lain2": "L",
    "koli_total": "M",
    "weight_kg": "N",
}
NUMERIC_FIELDS = ("serial_no", "tas", "krt", "krg", "lain2", "koli_total", "weight_kg")
TEXT_FIELDS = ("sender_name", "recipient_name", "address", "passport_no", "recipient_phone")
LAIN2 = ("hambal", "selimut", "drum", "kotak_besi")


# ─── Label dari workbook final ─────────────────────────────────────────


def load_labels(xlsx: Path, container: str) -> dict[str, dict[str, str]]:
    """Baris sheet Data Laut → {serial: {field: nilai}}. Header dicari lewat teks "Pengirim" di kolom B."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    name = next((n for n in wb.sheetnames if n.startswith(f"Data Laut Container-{container}")), None)
    if name is None:
        raise SystemExit(f'Sheet "Data Laut Container-{container}" tidak ditemukan di {xlsx.name}.')
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    col = {k: ord(v) - ord("A") for k, v in DATA_LAUT_COLUMNS.items()}
    header = next(
        (i for i, r in enumerate(rows) if len(r) > 1 and str(r[1] or "").strip().lower() == "pengirim"), None
    )
    if header is None:
        raise SystemExit('Header "Pengirim" di kolom B sheet Data Laut tidak ditemukan.')
    labels: dict[str, dict[str, str]] = {}
    for r in rows[header + 1 :]:
        serial = r[col["serial_no"]] if len(r) > col["serial_no"] else None
        if serial is None or not str(serial).strip().isdigit():
            continue
        labels[str(serial).strip()] = {
            k: ("" if i >= len(r) or r[i] is None else str(r[i]).strip()) for k, i in col.items()
        }
    return labels


# ─── Normalisasi hasil VLM ke bentuk kolom Excel ──────────────────────


def to_excel_shape(read: dict[str, Any]) -> dict[str, str]:
    """Output schema VLM → nilai yang sebanding dengan kolom Data Laut."""
    pk = read.get("packages") or {}

    def qty(key: str) -> int:
        digits = parse_count(str(pk.get(key) or ""))
        return int(digits) if digits else 0

    weight = parse_weight_kg(str(read.get("weight_kg") or ""))
    return {
        "serial_no": normalize_numeric(str(read.get("serial_no") or "")),
        "sender_name": str(read.get("sender_name") or ""),
        "recipient_name": str(read.get("recipient_name") or ""),
        "address": str(read.get("address") or ""),
        "passport_no": normalize_passport(str(read.get("passport_no") or "")),
        "recipient_phone": normalize_phones(str(read.get("recipient_phone") or "")),
        "tas": str(qty("koper")),
        "krt": str(qty("karton")),
        "krg": str(qty("karung")),
        "lain2": str(sum(qty(k) for k in LAIN2)),
        "koli_total": parse_count(str(read.get("koli_total") or "")),
        "weight_kg": "" if weight is None else f"{weight:g}",
    }


def _canon_text(v: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", v.upper())


def _canon_num(v: str) -> str:
    try:
        return f"{float(v.replace(',', '.')):g}" if v else "0"
    except ValueError:
        return normalize_numeric(v) or "0"


def field_correct(key: str, predicted: str, truth: str) -> bool:
    if key in NUMERIC_FIELDS:
        return _canon_num(predicted) == _canon_num(truth)
    if key == "passport_no":
        return normalize_passport(predicted) == normalize_passport(truth)
    if key == "recipient_phone":
        return normalize_phones(predicted) == normalize_phones(truth)
    return _canon_text(predicted) == _canon_text(truth)


def char_error_rate(predicted: str, truth: str) -> float:
    t = _canon_text(truth)
    return Levenshtein.distance(_canon_text(predicted), t) / max(len(t), 1)


# ─── Laporan ───────────────────────────────────────────────────────────


@dataclass
class SpikeResult:
    images: int = 0
    matched: int = 0
    unmatched_serial: int = 0
    failed: int = 0
    latencies: list[float] = field(default_factory=list)
    tokens_in: list[int] = field(default_factory=list)
    tokens_out: list[int] = field(default_factory=list)
    correct: dict[str, int] = field(default_factory=dict)
    cer: dict[str, list[float]] = field(default_factory=dict)

    def add(self, pred: dict[str, str], truth: dict[str, str]) -> None:
        self.matched += 1
        for k in DATA_LAUT_COLUMNS:
            if field_correct(k, pred.get(k, ""), truth.get(k, "")):
                self.correct[k] = self.correct.get(k, 0) + 1
            if k in TEXT_FIELDS:
                self.cer.setdefault(k, []).append(char_error_rate(pred.get(k, ""), truth.get(k, "")))

    def accuracy(self, keys: tuple[str, ...] | None = None) -> float:
        keys = keys or tuple(DATA_LAUT_COLUMNS)
        total = self.matched * len(keys)
        return sum(self.correct.get(k, 0) for k in keys) / total if total else 0.0

    def report(self, price_in: float | None, price_out: float | None) -> str:
        p95 = statistics.quantiles(self.latencies, n=20)[18] if len(self.latencies) >= 2 else 0.0
        lines = [
            "# Spike VLM",
            "",
            (
                f"Gambar: {self.images} · cocok dengan label: {self.matched} · serial tak cocok: "
                f"{self.unmatched_serial} · gagal: {self.failed}"
            ),
            "",
            "| Field | Benar | Akurasi | CER |",
            "| --- | --- | --- | --- |",
        ]
        for k in DATA_LAUT_COLUMNS:
            n = self.correct.get(k, 0)
            acc = n / self.matched if self.matched else 0
            cer = f"{statistics.mean(self.cer[k]):.1%}" if self.cer.get(k) else "–"
            lines.append(f"| {k} | {n}/{self.matched} | {acc:.1%} | {cer} |")
        overall, numeric = self.accuracy(), self.accuracy(NUMERIC_FIELDS)
        lines += [
            "",
            f"**Akurasi semua field: {overall:.1%}** (gate Fase 0: ≥ 80%)",
            f"**Akurasi field numerik: {numeric:.1%}** (gate Fase 0: ≥ 95%)",
            f"Waktu per resi: median {statistics.median(self.latencies or [0]):.1f} dtk · p95 {p95:.1f} dtk (target ≤ 12 dtk)",
        ]
        if self.tokens_in:
            ti, to = statistics.mean(self.tokens_in), statistics.mean(self.tokens_out)
            lines.append(f"Token rata-rata per resi: input {ti:.0f} · output {to:.0f}")
            if price_in is not None and price_out is not None:
                cost = (ti * price_in + to * price_out) / 1_000_000
                lines.append(
                    f"Biaya per resi ≈ USD {cost:.4f} · per container (250 resi) ≈ USD {cost * 250:.2f}"
                )
        passed = overall >= 0.80 and numeric >= 0.95 and p95 <= 12
        lines += ["", f"## Kesimpulan: {'LULUS' if passed else 'TIDAK LULUS'} gate Fase 0"]
        return "\n".join(lines)


# ─── Klien tiruan (tanpa HF_TOKEN) ────────────────────────────────────


class FakeClient:
    """Menjawab dari label berdasarkan urutan gambar — hanya untuk menguji alat ini."""

    def __init__(self, answers: list[dict[str, Any]]):
        self._answers = iter(answers)
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, **_kwargs: Any) -> Any:
        content = json.dumps(next(self._answers))
        usage = SimpleNamespace(prompt_tokens=1500, completion_tokens=300)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))], usage=usage
        )


def label_as_read(truth: dict[str, str]) -> dict[str, Any]:
    return {
        "serial_no": truth["serial_no"],
        "sender_name": truth["sender_name"].upper(),
        "recipient_name": truth["recipient_name"].upper(),
        "address": truth["address"].upper(),
        "passport_no": truth["passport_no"],
        "recipient_phone": truth["recipient_phone"],
        "packages": {
            "koper": truth["tas"] or None,
            "karton": truth["krt"] or None,
            "karung": truth["krg"] or None,
            "hambal": truth["lain2"] or None,
            "selimut": None,
            "drum": None,
            "kotak_besi": None,
        },
        "koli_total": truth["koli_total"],
        "weight_kg": truth["weight_kg"],
    }


# ─── Perintah ─────────────────────────────────────────────────────────


def cmd_pdf(args: argparse.Namespace) -> None:
    import pypdfium2 as pdfium

    out = Path(args.out).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(Path(args.pdf).expanduser()))
    for i, page in enumerate(pdf, start=1):
        page.render(scale=args.dpi / 72).to_pil().convert("RGB").save(out / f"hal-{i:03d}.jpg", quality=90)
    print(f"{len(pdf)} halaman → {out}")


def cmd_run(args: argparse.Namespace) -> None:
    images = sorted(
        p for p in Path(args.images).expanduser().iterdir() if p.suffix.lower() in {".jpg", ".jpeg", ".png"}
    )
    images = images[: args.n]
    labels = load_labels(Path(args.labels).expanduser(), args.container)
    settings = load_settings()
    if args.fake:
        client: Any = FakeClient([label_as_read(v) for v in list(labels.values())[: len(images)]])
    else:
        client = make_client(settings)

    out = Path(args.out).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    result, raw = SpikeResult(images=len(images)), []
    for img in images:
        t0 = time.perf_counter()
        try:
            resp = client.chat.completions.create(**_request(settings, img))
            read = json.loads(resp.choices[0].message.content)
        except Exception as exc:  # noqa: BLE001 — setiap kegagalan dicatat, spike tetap lanjut
            result.failed += 1
            raw.append({"image": img.name, "error": f"{type(exc).__name__}: {exc}"})
            continue
        result.latencies.append(time.perf_counter() - t0)
        if getattr(resp, "usage", None):
            result.tokens_in.append(resp.usage.prompt_tokens)
            result.tokens_out.append(resp.usage.completion_tokens)
        pred = to_excel_shape(read)
        raw.append({"image": img.name, "read": read})
        truth = labels.get(pred["serial_no"])
        if truth is None:
            result.unmatched_serial += 1
        else:
            result.add(pred, truth)

    (out / "bacaan-mentah.json").write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
    text = result.report(args.price_in, args.price_out)
    (out / "LAPORAN.md").write_text(text + "\n", encoding="utf-8")
    print(text)
    print(f"\nFile hasil (berisi data pribadi, jangan dibagikan): {out}")


def _request(settings: Any, img: Path) -> dict[str, Any]:
    from .hf_client import build_request

    return build_request(settings, [image_data_url(img)], INSTRUCTION)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="spike_vlm", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_pdf = sub.add_parser("pdf", help="pecah PDF menjadi gambar per halaman")
    p_pdf.add_argument("pdf")
    p_pdf.add_argument("--out", required=True)
    p_pdf.add_argument("--dpi", type=int, default=300)
    p_run = sub.add_parser("run", help="jalankan VLM dan bandingkan dengan label")
    p_run.add_argument("--images", required=True)
    p_run.add_argument("--labels", required=True, help="workbook final berisi sheet Data Laut Container-{no}")
    p_run.add_argument("--container", required=True)
    p_run.add_argument("--n", type=int, default=50)
    p_run.add_argument("--out", required=True)
    p_run.add_argument("--price-in", type=float, help="USD per 1 juta token input")
    p_run.add_argument("--price-out", type=float, help="USD per 1 juta token output")
    p_run.add_argument("--fake", action="store_true", help="klien tiruan, tanpa HF_TOKEN")
    args = parser.parse_args(argv)
    {"pdf": cmd_pdf, "run": cmd_run}[args.cmd](args)


if __name__ == "__main__":
    main()
